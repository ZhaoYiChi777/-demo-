import io
import os
import unittest
import tempfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
os.environ.setdefault('LLM_RETRIEVAL_ENABLED','false')
import main


class BriefImportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.paths=[patch.object(main,'DB_PATH',Path(self.temp.name)/'test.db'),patch.object(main,'DATA_DIR',Path(self.temp.name))]
        for item in self.paths:item.start()
        self.client=TestClient(main.app);self.client.__enter__()
        from test_accounts_support import login
        login(self.client)

    def tearDown(self):
        self.client.__exit__(None,None,None)
        for item in reversed(self.paths):item.stop()
        self.temp.cleanup()
    def upload(self,name,data):return self.client.post('/api/workbench/brief/import',files={'file':(name,data)})

    def test_chinese_text_and_markdown(self):
        for name,encoding in [('要求.txt','gb18030'),('要求.md','utf-8-sig')]:
            response=self.upload(name,'验收要求：提供性能试验记录。'.encode(encoding))
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.json()['text'],'验收要求：提供性能试验记录。')

    def test_docx_paragraph_and_table(self):
        document=Document();document.add_paragraph('任务目标：更新控制系统。')
        table=document.add_table(rows=1,cols=2);table.cell(0,0).text='验收要求';table.cell(0,1).text='提交试验记录'
        data=io.BytesIO();document.save(data)
        response=self.upload('任务书.docx',data.getvalue())
        self.assertEqual(response.status_code,200)
        self.assertIn('任务目标',response.json()['text'])
        self.assertIn('提交试验记录',response.json()['text'])

    def test_short_pdf_does_not_require_ocr(self):
        writer=PdfWriter();page=writer.add_blank_page(width=400,height=400)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 30 350 Td (Acceptance: submit test records.) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream)
        data=io.BytesIO();writer.write(data)
        with patch.object(main,'ocr_pdf',side_effect=AssertionError('Short text PDF must not use OCR')):
            response=self.upload('requirements.pdf',data.getvalue())
        self.assertEqual(response.status_code,200)
        self.assertIn('submit test records',response.json()['text'])

    def test_doc_converter_unavailable_is_actionable(self):
        with patch.object(main,'convert_doc_to_docx',side_effect=FileNotFoundError('libreoffice')):
            response=self.upload('任务.doc',b'legacy document')
        self.assertEqual(response.status_code,422)
        self.assertIn('另存为 .docx',response.json()['detail'])

    def test_empty_and_unsupported_do_not_return_requirements(self):
        for name,data in [('empty.txt',b''),('file.zip',b'zip')]:
            response=self.upload(name,data)
            self.assertEqual(response.status_code,422)
            self.assertNotIn('text',response.json())
