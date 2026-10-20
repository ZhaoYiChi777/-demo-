"""Rule snapshots, evidence-linked checks and reports for the review workbench."""
from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone
from typing import Literal
from uuid import uuid4

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field, model_validator
from urllib.parse import quote
from quality_scoring import quality_score
from semantic_judgment import SemanticJudge
from rule_contracts import contracted, excluded
from atomic_contracts import source_metadata
from review_conclusion import project_conclusion
from fsr_rubric import POLICY as FSR_POLICY, POLICIES as FSR_POLICIES, MAX_SCORE, COMMON, BRANCHES, select_branch, scoring_rules

CATEGORIES = {"common": "通用标准", "special": "专项标准", "topic": "课题要求"}
VERDICTS = {"pass": "符合", "fail": "不符合", "pending": "待确认", "not_applicable": "不适用"}


def now():
    return datetime.now(timezone.utc).isoformat()


def dumps(value):
    return json.dumps(value, ensure_ascii=False)


class RuleInput(BaseModel):
    name: str = Field(min_length=1, max_length=160)
    category: Literal["common", "special", "topic"]
    standard: str = Field(min_length=1)
    content: str = Field(min_length=1)
    enabled: bool = True
    mode: Literal["manual", "presence"] = "manual"
    terms: list[str] = Field(default_factory=list)
    suggestion: str = "请补充对应材料和核验依据。"
    specialty: str = "全部"
    project_id: str | None = None

    @model_validator(mode="after")
    def check_terms(self):
        self.name = self.name.strip()
        self.standard = self.standard.strip()
        self.content = self.content.strip()
        self.terms = list(dict.fromkeys(t.strip() for t in self.terms if t.strip()))
        if not self.name or not self.standard or not self.content:
            raise ValueError("规则名称、标准和内容不能为空")
        if self.mode == "presence" and not self.terms:
            raise ValueError("字段存在性检查至少需要一个核验字段")
        return self


class RunInput(BaseModel):
    material_package_id: int | None = None
    scoring_branch: Literal['safety','energy','heat','turbine','innovation'] | None = None


class BriefInput(BaseModel):
    text: str


class CheckInput(BaseModel):
    result: Literal["pass", "fail", "pending", "not_applicable"]
    reason: str = Field(min_length=1)
    suggestion: str = Field(min_length=1)


class Workbench:
    def __init__(self, db, knowledge):
        self.db = db
        self.knowledge = knowledge
        self.retriever = None
        self.access = None
        self.router = APIRouter(prefix="/api/workbench")
        self.router.add_api_route("/rules", self.rules, methods=["GET"])
        self.router.add_api_route("/rules", self.create_rule, methods=["POST"])
        self.router.add_api_route("/rules/{rule_id}", self.update_rule, methods=["PUT"])
        self.router.add_api_route("/projects/{project_id}/materials", self.materials, methods=["GET"])
        self.router.add_api_route("/projects/{project_id}/brief", self.save_brief, methods=["PUT"])
        self.router.add_api_route("/projects/{project_id}/runs", self.run, methods=["POST"])
        self.router.add_api_route("/projects/{project_id}/runs", self.history, methods=["GET"])
        self.router.add_api_route("/projects/{project_id}/runs/latest", self.latest, methods=["GET"])
        self.router.add_api_route("/runs/{run_id}", self.get_run, methods=["GET"])
        self.router.add_api_route("/runs/{run_id}/checks/{check_id}", self.review_check, methods=["PATCH"])
        self.router.add_api_route("/runs/{run_id}/report.docx", self.report, methods=["GET"])
        self.router.add_api_route("/runs/{run_id}/archive", self.archive, methods=["POST"])
        self.router.add_api_route("/runs/{run_id}/qualify", self.qualify, methods=["POST"])
        self.router.add_api_route("/demo", self.demo, methods=["POST"])
        self.router.add_api_route('/scoring-policy', self.scoring_policy, methods=['GET'])

    def scoring_policy(self):
        return dict(policy=FSR_POLICY,common=COMMON,branches=BRANCHES,
                    grading={'A':85,'B':75,'C':60,'D':0},maximum=MAX_SCORE,caps={'overall_maximum':MAX_SCORE,'special_chapter_only':49},
                    consistency_policy='关键数据矛盾按分项计分并列明问题，不自动封顶；项目结论按原规则等级汇总。')

    def init_tables(self):
        with self.db() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS wb_rules(id TEXT PRIMARY KEY, payload TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS wb_briefs(project_id TEXT PRIMARY KEY, text TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS wb_runs(id TEXT PRIMARY KEY, project_id TEXT NOT NULL, payload TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS wb_runs_project ON wb_runs(project_id);
            """)
            defaults = [
                ("GEN-001", "common", "项目基本信息字段", "材料需出现项目名称、建设单位、机组编号三个字段。", ["项目名称", "建设单位", "机组"], "presence"),
                ("GEN-002", "common", "申报章节完整性", "材料需包含建设必要性、技术方案、投资估算三个章节标识。", ["建设必要性", "技术方案", "投资估算"], "presence"),
                ("GEN-003", "common", "数据与逻辑一致性", "检查项目范围、设备数量、投资金额在章节之间是否一致。", ["投资", "机组", "设备"], "manual"),
                ("TECH-001", "special", "专业技术方案符合性", "核验设备配置、安全措施、系统切换与适用专业规范的符合性。", ["安全", "切换", "配置"], "manual"),
                ("TOPIC-001", "topic", "课题目标覆盖", "材料建设内容应逐项覆盖本次任务书的目标与交付成果。", ["目标", "成果"], "manual"),
                ("TOPIC-002", "topic", "课题验收指标", "逐项核验任务书目标值、验收方法及完成时限。", ["指标", "验收", "时限"], "manual"),
            ]
            for rid, cat, name, standard, terms, mode in defaults:
                rule = RuleInput(name=name, category=cat, standard=standard, content=standard, terms=terms, mode=mode).model_dump()
                rule.update(id=rid, origin="default", version=1, updated_at=now(), source="系统默认规则")
                conn.execute("INSERT OR IGNORE INTO wb_rules VALUES(?,?)", (rid, dumps(rule)))
            for row in self.knowledge:
                rid = row.get("规则ID")
                if not rid:
                    continue
                cat = "common" if row.get("类别") == "DOC" or any(t in row.get("类别", "") for t in ["完整", "格式", "基础", "资料"]) else "special"
                content = row.get("标准评审意见") or row.get("触发条件（自然语言）") or row.get("规则名称", rid)
                rule = RuleInput(name=row.get("规则名称") or rid, category=cat,
                                 standard=row.get("触发条件（自然语言）") or content, content=content,
                                 terms=re.split(r"[；;、，,\s]+", row.get("所需证据字段", "")),
                                 specialty=row.get("专业") or "汽机",
                                 suggestion=row.get("整改/补充要求") or "请按规则补充材料并核验。").model_dump()
                rule.update(id=rid, origin="default", version=1, updated_at=now(), source="现有原子规则库", source_metadata=source_metadata(row))
                conn.execute("INSERT OR IGNORE INTO wb_rules VALUES(?,?)", (rid, dumps(rule)))
                existing = json.loads(conn.execute("SELECT payload FROM wb_rules WHERE id=?", (rid,)).fetchone()["payload"])
                if existing["origin"] == "default" and (existing["category"] != cat or existing.get('source_metadata') != rule['source_metadata']):
                    existing.update(category=cat, source_metadata=rule['source_metadata'], updated_at=now())
                    conn.execute("UPDATE wb_rules SET payload=? WHERE id=?", (dumps(existing), rid))

    def project(self, project_id):
        if self.access:
            self.access.check_project(project_id)
        with self.db() as conn:
            row = conn.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
        if not row:
            raise HTTPException(404, "项目不存在")
        return dict(row)

    def rules(self, project_id: str | None = None, material_package_id: int | None = None):
        with self.db() as conn:
            rules = [json.loads(r["payload"]) for r in conn.execute("SELECT payload FROM wb_rules ORDER BY id")]
        if project_id:
            project = self.materials(project_id, material_package_id)['project'] if material_package_id is not None else self.project(project_id)
            demo = project_id == "DEMO-REVIEW"
            specialty = {"TURB": "汽机", "BOIL": "锅炉", "ELEC": "电气", "turbine": "汽机", "boiler": "锅炉", "electrical": "电气"}.get(project.get("specialty"), project.get("specialty"))
            rules = [r for r in rules if r.get("project_id") == project_id or
                     (not demo and not r.get("project_id") and (r["specialty"] == "全部" or not specialty or r["specialty"] == specialty))]
        else:
            rules = [r for r in rules if not r.get("project_id")]
        return {"rules": [contracted(r) for r in rules]}

    def create_rule(self, payload: RuleInput):
        if payload.project_id:
            self.project(payload.project_id)
        rule = payload.model_dump()
        rule.update(id=f"USR-{uuid4().hex[:10]}", origin="custom", version=1, updated_at=now(), source="用户自定义规则")
        with self.db() as conn:
            conn.execute("INSERT INTO wb_rules VALUES(?,?)", (rule["id"], dumps(rule)))
        return rule

    def update_rule(self, rule_id: str, payload: RuleInput):
        with self.db() as conn:
            row = conn.execute("SELECT payload FROM wb_rules WHERE id=?", (rule_id,)).fetchone()
            if not row:
                raise HTTPException(404, "规则不存在")
            old = json.loads(row["payload"])
            rule = {**old, **payload.model_dump(), "project_id": old.get("project_id"), "origin": "custom",
                    "version": old["version"] + 1, "updated_at": now()}
            for key in ('criteria', 'contract_version', 'review_guidance', 'applicability_condition', 'source_metadata'):
                rule.pop(key, None)
            conn.execute("UPDATE wb_rules SET payload=? WHERE id=?", (dumps(rule), rule_id))
        return rule

    def materials(self, project_id: str, material_package_id: int | None = None):
        project = self.project(project_id)
        with self.db() as conn:
            if material_package_id is not None and not conn.execute("SELECT 1 FROM material_packages WHERE id=? AND project_id=?", (material_package_id, project_id)).fetchone():
                raise HTTPException(404, "资料包不存在")
            params = [project_id]
            clause = ""
            if material_package_id is not None:
                clause = " AND EXISTS(SELECT 1 FROM project_documents pd WHERE pd.source_id=s.id AND pd.project_id=s.project_id AND pd.material_package_id=?)"
                params.append(material_package_id)
            sources = [dict(r) for r in conn.execute("SELECT s.*, pd.material_package_id FROM document_sources s LEFT JOIN project_documents pd ON pd.source_id=s.id AND pd.project_id=s.project_id WHERE s.project_id=?" + clause + " ORDER BY s.id", params)]
            doc_ids = list({s["document_id"] for s in sources if s["document_category"] in ("project_report", "drawing") and s["document_id"]})
            chunks = []
            if doc_ids:
                chunks = [dict(r) for r in conn.execute("SELECT c.id, c.document_id, c.section_path, c.page, c.text, d.filename FROM chunks c JOIN documents d ON c.document_id=d.id WHERE c.project_id=? AND c.document_id IN (" + ",".join("?" for _ in doc_ids) + ") ORDER BY c.document_id,c.id", [project_id, *doc_ids])]
            if material_package_id is not None:
                # A grouped report uses its own document specialty, never the last upload's specialty.
                specialties = {}
                for assessment in conn.execute('SELECT document_id,primary_specialty FROM specialty_assessments WHERE project_id=? ORDER BY id DESC', (project_id,)):
                    if assessment['document_id'] in doc_ids:
                        specialties.setdefault(assessment['document_id'], assessment['primary_specialty'])
                unique = {s for s in specialties.values() if s}
                project['specialty'] = next(iter(unique)) if len(unique) == 1 else None
            brief = conn.execute("SELECT text FROM wb_briefs WHERE project_id=?", (project_id,)).fetchone()
            package = conn.execute("SELECT name FROM material_packages WHERE id=? AND project_id=?", (material_package_id, project_id)).fetchone() if material_package_id else None
        return {"project": project, "sources": sources, "chunks": chunks, "brief": brief["text"] if brief else "", "material_package_id": material_package_id, "scope_name": package["name"] if package else "整个项目"}

    def save_brief(self, project_id: str, payload: BriefInput):
        self.project(project_id)
        with self.db() as conn:
            old = conn.execute('SELECT text FROM wb_briefs WHERE project_id=?', (project_id,)).fetchone()
            conn.execute("INSERT INTO wb_briefs VALUES(?,?) ON CONFLICT(project_id) DO UPDATE SET text=excluded.text", (project_id, payload.text.strip()))
        if self.access:
            self.access.log('materials', '更新任务要求', self.project(project_id)['name'], project_id,
                            detail={'before': old['text'] if old else '', 'after': payload.text.strip()})
        return {"text": payload.text.strip()}

    def run(self, project_id: str, payload: RunInput):
        material = self.materials(project_id, payload.material_package_id)
        if not material["chunks"]:
            raise HTTPException(409, "当前范围没有可评审的原文，请上传申报材料并确认解析及资料分类。")
        rules = [r for r in self.rules(project_id, payload.material_package_id)["rules"] if r["enabled"]]
        branch = select_branch(material)
        if payload.scoring_branch:
            branch = dict(id=payload.scoring_branch,name=BRANCHES[payload.scoring_branch][0],reason='用户选择的本次项目分支',source='user')
        if project_id != 'DEMO-REVIEW':
            rules += scoring_rules(branch)
        retrieved, retrieval_status = self.retriever.retrieve(material, rules) if self.retriever else ({}, {'status':'disabled'})
        checks = []
        full_text = "\n".join(c["text"] for c in material["chunks"])
        incomplete = any(s["parse_status"] != "parsed" for s in material["sources"] if s["document_category"] in ("project_report", "drawing"))
        for rule in rules:
            terms = rule["terms"]
            missing = [t for t in terms if t not in full_text]
            if rule['id'] in retrieved:
                evidence = retrieved[rule['id']]['evidence']
            else:
                ranked = sorted(material['chunks'], key=lambda c:sum(t in c['text'] for t in terms), reverse=True)
                evidence = [c for c in ranked if any(t in c['text'] for t in terms)][:3]
            result = "pending"
            reason = "已检索相关原文；此项涉及专业或语义判断，需要按标准逐项复核。"
            if rule["mode"] == "presence":
                if missing:
                    reason = f"在本次已解析材料中未检出字段：{'、'.join(missing)}。"
                    result = "pending"
                    if incomplete:
                        reason += "存在未完成解析的材料，暂无法确认缺失。"
                else:
                    result = "pass"
                    reason = "已检出全部指定字段；此项仅判断字段出现情况，数据正确性由独立检查项核验。"
            if rule["category"] == "topic" and not material["brief"]:
                result, reason = "pending", "尚未录入本次课题任务书或验收要求，缺少比较基准。"
            if rule['mode']=='manual':
                reason = ('已检索到相关原文；' if evidence else '尚未检索到可核验的原文；') + '此项涉及专业或语义判断，需要按标准逐项复核。'
                if rule['category']=='topic' and not material['brief']:
                    reason = '尚未录入本次课题任务书或验收要求，缺少比较基准。'
            details = retrieved.get(rule['id'],{}).get('retrieval',retrieval_status)
            label = {'complete':'字段匹配 + 大模型语义检索', 'partial':'字段匹配 + 部分语义检索',
                     'unavailable':'字段匹配（语义检索未完成）'}.get(retrieval_status['status'],'字段匹配（语义检索未启用）')
            checks.append({"id": rule["id"], "rule": rule, "category": rule["category"], "result": result,
                           "reason": reason, "suggestion": rule["suggestion"] if result != "pass" else "本项无需修改，保留现有材料。",
                           "evidence": evidence, "missing": missing, "retrieval": details, "method": label + ('；字段存在性检查' if rule['mode']=='presence' else '；按判据评审'),
                           "reviewed_by": None})
        judgments, judgment_status = SemanticJudge(self.retriever.gateway, self.retriever.enabled).evaluate(material, checks) if self.retriever else ({}, {'status':'disabled'})
        for check in checks:
            decision = judgments.get(check['id'])
            if not decision:
                continue
            check['judgment'] = decision
            check['result'] = decision['result']
            check['reason'] = decision['reason']
            if decision['status'] in ('completed','partial'):
                check['method'] += '；语义等价核验' if check['rule']['mode'] == 'presence' else '；按判据语义评审'
                check['reason'] += '\n' + '\n'.join(f"{c['criterion']}：{VERDICTS[c['result']]}。{c['reason']}" for c in decision['criteria'])
                # Cross-rule evidence from the same review scope remains navigable and exportable.
                for criterion in decision['criteria']:
                    for ref in criterion.get('evidence', []):
                        if not any(e['id'] == ref['paragraph_id'] for e in check['evidence']):
                            original = next(c for c in material['chunks'] if c['id'] == ref['paragraph_id'])
                            check['evidence'].append({**original, 'retrieval_sources':['judgment'],
                                'semantic_quotes':[{'quote':ref['quote'], 'reason':criterion['reason']}]})
            else:
                check['method'] += '；语义判定未完成'
            check['suggestion'] = (decision.get('suggestion') or '本项无需修改，保留现有材料与核验依据。' if decision['result'] == 'pass' else
                                   decision.get('suggestion') or check['rule']['suggestion'])
        run = {"id": f"REV-{uuid4().hex[:12]}", "project_id": project_id, "created_at": now(), "updated_at": now(),
               "status": "completed", "is_demo": project_id == "DEMO-REVIEW", "material_package_id": payload.material_package_id,
               "project": material["project"], "sources": material["sources"], "chunks": material["chunks"], "brief": material["brief"], "checks": checks, "scope_name": material["scope_name"], "retrieval": retrieval_status, "judgment": judgment_status}
        if self.access:
            from accounts import actor
            if actor.get():
                run['created_by'] = actor.get()['username']
        if project_id != 'DEMO-REVIEW':
            run.update(scoring_policy=FSR_POLICY,scoring_branch=branch)
        self.aggregate(run)
        self.save_run(run)
        return run

    @staticmethod
    def aggregate(run):
        checks = run["checks"]
        def verdict(items):
            if items and all(i.get('result') == 'not_applicable' or ('rule' in i and excluded(i)) for i in items):
                return 'not_applicable'
            items = [i for i in items if i.get('result') != 'not_applicable' and not ('rule' in i and excluded(i))]
            if any(i["result"] == "fail" for i in items):
                return "fail"
            if not items or any(i["result"] == "pending" for i in items):
                return "pending"
            return "pass"
        for c in checks:
            if excluded(c):
                c['result'] = 'not_applicable'
        run["counts"] = {v: sum(c["result"] == v for c in checks) for v in VERDICTS}
        run["counts"]["total"] = len(checks)
        run["categories"] = [{"id": key, "name": name, "result": verdict([c for c in checks if c["category"] == key]),
                              "total": sum(c["category"] == key for c in checks),
                              "issues": sum(c["category"] == key and c["result"] in ('fail','pending') for c in checks)} for key, name in CATEGORIES.items()]
        run["result"] = verdict(run["categories"])
        if run['result'] == 'not_applicable':
            run['result'] = 'pending'
        run["topic_result"] = next(c["result"] for c in run["categories"] if c["id"] == "topic")
        run['quality_score'] = quality_score(run)
        if any(c['rule'].get('contract_version') == 4 for c in checks):
            run['project_conclusion'] = project_conclusion(checks)
        run["recommendation"] = {"pass": "符合当前启用的评审标准，可以进入下一阶段。", "fail": "建议按评审意见修改材料后重新提交。", "pending": "请补充待确认事项的证据并完成核验，再确定是否进入下一阶段。"}[run["result"]]
        if run.get('project_conclusion'):
            run['recommendation'] = run['project_conclusion']['explanation']
        run["summary"] = f"本次共检查 {len(checks)} 项规则，其中 {run['counts']['pass']} 项符合、{run['counts']['fail']} 项不符合、{run['counts']['pending']} 项待确认。"
        if run['counts']['not_applicable']:
            run['summary'] += f"另有 {run['counts']['not_applicable']} 项不适用，已退出计分与符合性判定。"

    def save_run(self, run):
        with self.db() as conn:
            previous = conn.execute('SELECT payload FROM wb_runs WHERE id=?', (run['id'],)).fetchone()
            conn.execute("INSERT INTO wb_runs VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET payload=excluded.payload", (run["id"], run["project_id"], dumps(run)))
        if self.access:
            title = run.get('scope_name') if run.get('scope_name') != '整个项目' else run['project']['name']
            if not previous:
                self.access.log('reviews', '完成评审', title or run['project']['name'], run['project_id'], run['id'],
                                detail={'result': run['result']}, event_key='run:' + run['id'], at=run['created_at'])
            else:
                old = json.loads(previous['payload'])
                if old.get('qualification') != run.get('qualification'):
                    action = '人工确认合格'
                elif old['status'] != run['status']:
                    action = '归档报告'
                elif old.get('checks') != run.get('checks'):
                    action = '复核检查项'
                else:
                    return
                changes = [{'name': c['rule']['name'], 'before': before, 'after': c}
                           for before, c in zip(old.get('checks', []), run.get('checks', [])) if before != c]
                self.access.log('manual', action, title or run['project']['name'], run['project_id'], run['id'],
                                detail={'changes': changes, 'before': {'status': old['status'], 'result': old['result']},
                                        'after': {'status': run['status'], 'result': run['result'], 'qualification': run.get('qualification')}})

    def get_run(self, run_id: str):
        with self.db() as conn:
            row = conn.execute("SELECT payload FROM wb_runs WHERE id=?", (run_id,)).fetchone()
        if not row:
            raise HTTPException(404, "评审记录不存在")
        run = json.loads(row["payload"])
        if self.access:
            self.access.check_project(run['project_id'])
        return run

    def latest(self, project_id: str):
        self.project(project_id)
        with self.db() as conn:
            row = conn.execute("SELECT payload FROM wb_runs WHERE project_id=? ORDER BY rowid DESC LIMIT 1", (project_id,)).fetchone()
        return json.loads(row["payload"]) if row else None

    def history(self, project_id: str):
        self.project(project_id)
        with self.db() as conn:
            runs = [json.loads(row["payload"]) for row in conn.execute("SELECT payload FROM wb_runs WHERE project_id=? ORDER BY rowid DESC", (project_id,))]
        return {"runs": [{key: run.get(key) for key in ("id", "created_at", "status", "result", "scope_name", "qualification", "project_conclusion")} for run in runs]}

    def review_check(self, run_id: str, check_id: str, payload: CheckInput):
        run = self.get_run(run_id)
        if run["status"] == "archived":
            raise HTTPException(409, "已归档评审不可修改，请重新评审。")
        check = next((c for c in run["checks"] if c["id"] == check_id), None)
        if not check:
            raise HTTPException(404, "检查项不存在")
        if not payload.reason.strip() or not payload.suggestion.strip():
            raise HTTPException(422, "请填写判断原因和修改建议。")
        check.setdefault("history", []).append({"result": check["result"], "reason": check["reason"], "suggestion": check["suggestion"], "at": now()})
        from accounts import actor
        check.update(payload.model_dump(), reviewed_by=actor.get()['username'] if actor.get() else '人工复核', reviewed_at=now())
        run["updated_at"] = now()
        self.aggregate(run)
        self.save_run(run)
        return run

    def archive(self, run_id: str):
        run = self.get_run(run_id)
        run.update(status="archived", archived_at=now())
        self.save_run(run)
        return run

    def qualify(self, run_id: str):
        """Retain the current snapshot as manually qualified; automated findings stay intact."""
        run = self.get_run(run_id)
        if run.get('qualification'):
            return run
        timestamp = now()
        run.update(qualification={'result': 'qualified', 'method': 'manual', 'confirmed_at': timestamp,
                                  'original_result': run['result']},
                   status='archived', archived_at=run.get('archived_at') or timestamp)
        self.save_run(run)
        return run

    def report(self, run_id: str):
        run = self.get_run(run_id)
        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(10)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        document.add_heading("项目评审报告", 0)
        if run["is_demo"]:
            document.add_paragraph("演示样例 · 本报告使用示例材料，仅用于功能体验。")
        document.add_heading("1. 评审基本信息", 1)
        for label, value in [("项目名称", run["project"]["name"]), ("项目编号", run["project_id"]), ("评审编号", run["id"]), ("评审时间", run["created_at"]), ("报告状态", "已归档" if run["status"] == "archived" else "当前评审快照")]:
            document.add_paragraph(f"{label}：{value}")
        for source in run["sources"]:
            document.add_paragraph(f"材料：{source['filename']}；解析状态：{source['parse_status']}")
        document.add_paragraph(f"评审范围：{run.get('scope_name') or ('资料包 ' + str(run['material_package_id']) if run['material_package_id'] else '整个项目')}")
        document.add_heading("2. 综合评审结论", 1)
        if run.get('qualification'):
            document.add_paragraph(f"人工确认结论：合格。确认时间：{run['qualification']['confirmed_at']}。")
            document.add_paragraph('本报告经人工确认作为合格报告保留；下列自动评审结果、问题及评分按确认时的原始记录保留。')
        conclusion_label = '原始自动评审结论' if run.get('qualification') else '综合结论'
        opinion = run.get('project_conclusion')
        document.add_paragraph(f"{conclusion_label}：{opinion['label'] if opinion else VERDICTS[run['result']]}。课题要求符合性：{VERDICTS[run['topic_result']]}。")
        if opinion:
            document.add_paragraph(opinion['explanation'])
            document.add_paragraph('逐项符合性：' + VERDICTS[run['result']] + '；与项目四级结论分别记录。')
        document.add_heading("3. 评审摘要", 1)
        document.add_paragraph(run["summary"])
        document.add_paragraph(opinion['method'] if opinion else "结论按不符合优先、待确认次之、全部符合才符合的顺序聚合。未启用的规则不计入检查数；某类别无检查项时为待确认。")
        document.add_heading("4. 三类评审结果", 1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, ["评审类别", "结论", "检查数", "问题数（含待确认）"]):
            cell.text = value
        for cat in run["categories"]:
            for cell, value in zip(table.add_row().cells, [cat["name"], VERDICTS[cat["result"]], str(cat["total"]), str(cat["issues"])]):
                cell.text = value
        document.add_heading("5. 主要问题", 1)
        issues = [c for c in run["checks"] if c["result"] != "pass"]
        for c in issues:
            document.add_paragraph(f"[{VERDICTS[c['result']]}] {c['rule']['name']}：{c['reason']}")
        if not issues:
            document.add_paragraph("本次检查未发现不符合或待确认项。")
        document.add_heading("6. 评审依据", 1)
        document.add_paragraph("课题要求：" + (run["brief"] or "尚未提供"))
        for c in run["checks"]:
            r = c["rule"]
            document.add_paragraph(f"{r['id']} · v{r['version']} · {r['name']} · {'系统默认' if r['origin'] == 'default' else '用户自定义'}\n判断标准：{r['standard']}\n规则内容：{r['content']}")
        document.add_heading("7. 详细评审意见", 1)
        for c in run["checks"]:
            document.add_heading(c["rule"]["name"], 2)
            document.add_paragraph(f"规则：{c['id']} v{c['rule']['version']}；分类：{CATEGORIES[c['category']]}；结果：{VERDICTS[c['result']]}")
            document.add_paragraph(f"执行方式：{c['method']}；核验状态：{c.get('reviewed_by') or '系统检查记录'}")
            document.add_paragraph("评审标准：" + c["rule"]["standard"])
            meta = c['rule'].get('source_metadata')
            if meta:
                document.add_paragraph(f"原规则范围：{meta['project_types']} / {meta['system_scope']}；适用条件：{c['rule'].get('applicability_condition','')}；问题级别：{meta['problem_level']}；结论影响：{meta['conclusion_impact']}")
            for criterion in c.get('judgment', {}).get('criteria', []):
                document.add_paragraph(f"自动判据：{criterion['criterion']}；{VERDICTS[criterion['result']]}。{criterion['reason']}")
                for ref in criterion.get('evidence', []):
                    document.add_paragraph(f"判定引用 · 片段 {ref['paragraph_id']}：{ref['quote']}")
                    if ref.get('object'):
                        state = {'before':'改造前','after':'改造后','scope':'本次范围','unspecified':'原文未明确'}.get(ref.get('state'),'原文未明确')
                        document.add_paragraph(f"对应对象：{ref['object']}；状态：{state}；数值与单位：{ref.get('value','')}{ref.get('unit','')}")
                if criterion.get('arithmetic'):
                    document.add_paragraph(criterion['arithmetic']['reason'])
            supplement = c.get('judgment',{}).get('supplemental_search')
            if supplement:
                document.add_paragraph('补检索判据：'+'；'.join(supplement['criteria']))
                document.add_paragraph(f"补检索完成：{supplement['retrieval']['completed_batches']}/{supplement['retrieval']['total_batches']}组")
            for window in c.get('retrieval',{}).get('failed_windows',[]):
                document.add_paragraph(f"未完成窗口：段落{window['paragraph_id']}，字符{window['start']}—{window['end']}；{window['error']}")
            if c["evidence"]:
                for evidence in c["evidence"]:
                    if evidence.get('retrieval_sources'):
                        document.add_paragraph('检索来源：'+' + '.join('语义命中' if s=='semantic' else '字段命中' for s in evidence['retrieval_sources']))
                    document.add_paragraph(f"原文：{evidence['filename']} / {evidence['section_path']} / 片段 {evidence['id']}\n{evidence['text']}")
                    for detail in evidence.get('semantic_quotes',[]):
                        document.add_paragraph('语义检索引文：'+detail['quote']+'；关联说明：'+detail['reason'])
            else:
                document.add_paragraph("原文证据：未定位到匹配片段，请查看本次材料全文；缺失项没有对应原句。")
            document.add_paragraph("原因：" + c["reason"])
            document.add_paragraph("修改建议：" + c["suggestion"])
        document.add_heading("8. 最终建议", 1)
        document.add_paragraph(run["recommendation"])
        quality = run.get('quality_score',{})
        if quality.get('policy') in FSR_POLICIES:
            document.add_heading('9. 报告成熟度分项评分', 1)
            document.add_paragraph(f"评分口径：{quality['label']}；分支：{quality['branch']['name']}。")
            score_label = '暂计分项得分' if quality.get('provisional') else '已确认得分'
            document.add_paragraph(f"{score_label}：{quality['total']}；可得分区间：{quality.get('lower_bound',quality['total'])}—{quality['upper_bound']}；核验完成度：{quality['coverage']}%；等级：{quality['grade'] or '待核验完成后确定'}。")
            document.add_paragraph(quality['explanation'])
            if 'processing_pending_points' in quality:
                document.add_paragraph(f"模型处理未完成：{quality['processing_pending_points']:.2f}分；材料或适用性待确认：{quality['evidence_pending_points']:.2f}分。" + ('本次暂计分不参与有效评分比较。' if not quality.get('comparable') else ''))
            for dimension in quality['dimensions']:
                document.add_heading(f"{dimension['name']}：{dimension['contribution']:.2f}/{dimension['effective_weight']:.2f} 分",2)
                if dimension['excluded']:
                    document.add_paragraph('不适用，已退出计分。')
                for item in dimension['items']:
                    document.add_paragraph(f"{item['name']}：{VERDICTS[item['result']]}。{item['reason']}")
            for cap in quality['caps']:
                document.add_paragraph(f"最高{cap['limit']}分：{cap['reason']}")
            for note in quality.get('decision_notes', []):
                document.add_paragraph('关键数据核验提醒：' + note['reason'] + '\n' + note['explanation'])
        output = io.BytesIO()
        document.save(output)
        filename = quote(f"评审报告-{run['id']}.docx")
        return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})

    def demo(self):
        project_id = "DEMO-REVIEW"
        with self.db() as conn:
            exists = conn.execute("SELECT 1 FROM projects WHERE id=?", (project_id,)).fetchone()
            if not exists:
                conn.execute("INSERT INTO projects(id,name,specialty,project_type,status,created_at) VALUES(?,?,?,?,?,?)", (project_id, "12、13号机组抗燃油系统改造", "汽机", "技术改造", "uploaded", now()))
                sections = [
                    ("一、项目概况", "项目名称：12、13号机组抗燃油系统改造。建设单位：示例热电厂。建设必要性：现有管路运行时间较长，拟更换油箱及管路，改善供油可靠性。"),
                    ("二、技术方案", "技术方案：更换抗燃油箱、泵组及配套管路。保留现有控制逻辑。施工期间采用分机组停机方式，调试后恢复供油。油箱区沿用现有布置。"),
                    ("三、投资估算", "投资估算：项目总投资 280 万元，其中设备费用 210 万元，施工及其他费用 70 万元。设备采购以最终招标结果为准。"),
                    ("四、项目目标与验收", "项目目标：改善供油稳定性，降低故障停机次数。交付成果：完成设备安装与调试，提供竣工图纸。计划在年度检修期间完成，验收由生产部门组织。"),
                ]
                doc = conn.execute("INSERT INTO documents(project_id,filename,document_type,file_hash,content,parse_status,created_at) VALUES(?,?,?,?,?,?,?)", (project_id, "抗燃油系统改造申报书（示例）.txt", "project_material", "demo-review-material", "\n\n".join(t for _, t in sections), "parsed", now())).lastrowid
                conn.execute("INSERT INTO document_sources(project_id,document_id,source_path,filename,file_hash,file_size,file_kind,parse_status,document_category,parse_method,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)", (project_id, doc, "演示材料", "抗燃油系统改造申报书（示例）.txt", "demo-review-material", 0, "txt", "parsed", "project_report", "direct", now()))
                for title, text in sections:
                    conn.execute("INSERT INTO chunks(project_id,document_id,chunk_type,section_path,text,created_at) VALUES(?,?,?,?,?,?)", (project_id, doc, "project_material", title, text, now()))
                conn.execute("INSERT INTO wb_briefs VALUES(?,?)", (project_id, "本课题要求完成抗燃油系统改造；材料需包含项目目标、交付成果、可靠性量化目标及验收方法。示例文字检查用于展示，技术适用性仍需专家确认。"))
                demo_rules = [
                    ("common", "项目基本信息完整性", ["项目名称", "建设单位", "机组"], "presence", "项目名称、建设单位、机组编号字段应齐全。", "补充项目基本信息。"),
                    ("common", "申报章节完整性", ["建设必要性", "技术方案", "投资估算"], "presence", "包含建设必要性、技术方案、投资估算章节。", "补充缺失章节。"),
                    ("common", "投资金额逻辑一致性", ["投资", "费用"], "manual", "核验分项费用之和与总投资一致，并复核计价口径。", "请复核 210 万元与 70 万元的计价依据和费用边界。"),
                    ("special", "油箱区安全措施说明", ["防火分隔", "泄漏收集"], "presence", "本示例规则要求材料显式列出防火分隔、泄漏收集两项说明。", "在技术方案章节补充防火分隔、泄漏收集措施及设计依据。"),
                    ("special", "施工组织说明", ["停机", "调试"], "presence", "材料应列出停机及调试安排；本项只检查相关字段。", "补充停机与调试安排。"),
                    ("topic", "任务目标与成果覆盖", ["项目目标", "交付成果"], "presence", "课题要求申报书明确列出项目目标和交付成果。", "在项目目标章节列出项目目标和交付成果。"),
                    ("topic", "量化目标与验收方法", ["可靠性量化目标", "验收方法"], "presence", "本示例任务书要求显式列出可靠性量化目标和验收方法。", "补充可靠性目标值、测量条件及验收方法。"),
                ]
                for i, (cat, name, terms, mode, standard, suggestion) in enumerate(demo_rules, 1):
                    rid = f"DEMO-{i:03}"
                    rule = RuleInput(name=name, category=cat, standard=standard, content=standard, terms=terms, mode=mode, suggestion=suggestion, project_id=project_id).model_dump()
                    rule.update(id=rid, origin="default", version=1, updated_at=now(), source="功能演示示例规则")
                    conn.execute("INSERT OR IGNORE INTO wb_rules VALUES(?,?)", (rid, dumps(rule)))
        return {"project_id": project_id, "run": self.latest(project_id) or self.run(project_id, RunInput())}
