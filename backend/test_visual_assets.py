import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject

import main
from visual_assets import Collector, scan_docx, scan_pdf


def image_bytes():
    image=Image.new('RGB',(640,360),'white')
    draw=ImageDraw.Draw(image)
    draw.line((60,40,60,300,580,300),fill='black',width=3)
    draw.line((60,260,180,220,320,130,560,80),fill='blue',width=5)
    draw.text((80,15),'Figure 1 - performance trend',fill='black')
    output=io.BytesIO();image.save(output,format='PNG');return output.getvalue()


def docx_bytes(image_only=False):
    doc=Document()
    if not image_only:
        doc.add_paragraph('项目名称：图表提取测试。图1 性能曲线')
    doc.add_picture(io.BytesIO(image_bytes()))
    if not image_only:
        doc.add_paragraph('表1 投资估算，单位：万元')
        table=doc.add_table(rows=3,cols=2)
        for row,values in zip(table.rows,[['项目','金额'],['设备费','100'],['合计','100']]):
            for cell,value in zip(row.cells,values):cell.text=value
        doc.add_paragraph('表后说明：费用按本项目范围统计。')
    out=io.BytesIO();doc.save(out);return out.getvalue()


def vector_pdf():
    writer=PdfWriter();page=writer.add_blank_page(width=400,height=300)
    stream=DecodedStreamObject()
    stream.set_data(b'0 0 0 RG 20 20 360 240 re 20 80 m 380 80 l 20 140 m 380 140 l 180 20 m 180 260 l S')
    page[NameObject('/Contents')]=writer._add_object(stream)
    writer.add_blank_page(width=400,height=300)
    out=io.BytesIO();writer.write(out);return out.getvalue()


class VisualExtractorTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()

    def test_docx_table_image_locations_units_and_context(self):
        collector=Collector(self.folder);scan_docx(docx_bytes(),collector)
        self.assertEqual([i['kind'] for i in collector.items],['image','table'])
        image,table=collector.items
        self.assertIn('图1',image['context'])
        self.assertIsNone(table['location']['page'])
        self.assertIn('单位：万元',table['context'])
        self.assertIn('表后说明',table['context'])
        self.assertEqual(table['rows'][1],['设备费','100'])
        self.assertFalse(table['needs_recognition'])
        self.assertTrue(image['needs_recognition'])
        self.assertTrue((self.folder/image['files']['preview']).exists())

    def test_pdf_vector_page_saved_with_correct_page_and_blank_page_skipped(self):
        collector=Collector(self.folder);scan_pdf(vector_pdf(),collector)
        self.assertFalse(collector.errors)
        self.assertEqual(len(collector.items),1)
        self.assertEqual(collector.items[0]['location']['page'],1)
        self.assertIn('vector_graphics_or_table',collector.items[0]['reasons'])
        with Image.open(self.folder/collector.items[0]['files']['preview']) as im:
            self.assertEqual(im.size,(800,600))

    def test_scanned_pdf_keeps_visual_even_without_text(self):
        image=Image.open(io.BytesIO(image_bytes()));out=io.BytesIO();image.save(out,format='PDF')
        collector=Collector(self.folder);scan_pdf(out.getvalue(),collector)
        self.assertEqual(len(collector.items),1)
        self.assertEqual(collector.items[0]['context'],'')
        self.assertTrue(collector.items[0]['needs_recognition'])

    def test_native_formula_saved_for_later_interpretation(self):
        doc=Document();p=doc.add_paragraph('公式说明')
        math=OxmlElement('m:oMath');run=OxmlElement('m:r');text=OxmlElement('m:t');text.text='x=1'
        run.append(text);math.append(run);p._p.append(math)
        out=io.BytesIO();doc.save(out)
        collector=Collector(self.folder);scan_docx(out.getvalue(),collector)
        self.assertEqual(collector.items[0]['kind'],'formula')
        self.assertTrue(collector.items[0]['needs_recognition'])

    def test_table_inside_content_control_keeps_rows_and_block(self):
        doc=Document();table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='设备费';table.cell(0,1).text='100万元'
        wrapper=OxmlElement('w:sdt');content=OxmlElement('w:sdtContent')
        caption=doc.add_paragraph('表1 投资估算，单位：万元')
        content.append(caption._p);content.append(table._tbl);wrapper.append(content);doc.element.body.insert(0,wrapper)
        out=io.BytesIO();doc.save(out)
        collector=Collector(self.folder);scan_docx(out.getvalue(),collector)
        self.assertEqual(len(collector.items),1)
        self.assertEqual(collector.items[0]['rows'],[['设备费','100万元']])
        self.assertEqual(collector.items[0]['location']['block_index'],1)
        self.assertIn('单位：万元',collector.items[0]['context'])


class VisualUploadTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();root=Path(self.temp.name)
        self.patches=[patch.object(main,'DB_PATH',root/'test.db'),patch.object(main,'DATA_DIR',root),
                      patch.object(main,'UPLOAD_DIR',root/'uploads'),patch.object(main,'CONVERTED_DIR',root/'converted')]
        for p in self.patches:p.start()
        self.client=TestClient(main.app);self.client.__enter__()
        from test_accounts_support import login
        login(self.client)
        self.project=self.client.post('/api/projects',data={'name':'图表上传测试'}).json()['project_id']

    def tearDown(self):
        self.client.__exit__(None,None,None)
        for p in reversed(self.patches):p.stop()
        self.temp.cleanup()

    def upload(self,name,data):
        response=self.client.post(f'/api/projects/{self.project}/documents',files={'file':(name,data)})
        self.assertEqual(response.status_code,200,response.text)
        source=self.client.get(f'/api/projects/{self.project}/sources').json()['source_files'][-1]
        assets=self.client.get(f"/api/projects/{self.project}/sources/{source['id']}/visual-assets").json()
        return response.json(),source,assets

    def test_upload_saves_downloadable_assets_without_ocr_and_keeps_review_text_separate(self):
        with patch('local_ocr.image_text',side_effect=AssertionError('No new OCR call expected')):
            result,source,assets=self.upload('报告.docx',docx_bytes())
        self.assertEqual(result['visual_assets']['count'],2)
        self.assertEqual(source['asset_scan_status'],'complete')
        table=next(a for a in assets['items'] if a['kind']=='table')
        self.assertIn('VIS-',table['citation'])
        self.assertEqual(self.client.get(table['files']['table']).json()['rows'][1],['设备费','100'])
        original=self.client.get(table['source_url'])
        self.assertEqual(original.status_code,200)
        self.assertTrue(original.content.startswith(b'PK'))
        with main.db() as conn:
            chunks=conn.execute('SELECT text FROM chunks WHERE project_id=?',(self.project,)).fetchall()
            self.assertTrue(chunks)
            self.assertTrue(all('VIS-' not in c['text'] for c in chunks))
        wrong=table['files']['table'].replace(self.project,'wrong-project')
        self.assertEqual(self.client.get(wrong).status_code,403)

    def test_image_only_docx_is_saved_even_when_no_text_document_created(self):
        _,source,assets=self.upload('只有图.docx',docx_bytes(image_only=True))
        self.assertEqual(source['parse_status'],'empty_content')
        self.assertIsNone(source['document_id'])
        self.assertEqual(assets['status'],'complete')
        self.assertEqual(len(assets['items']),1)

    def test_archive_leaf_has_own_identity_and_original(self):
        original=docx_bytes();archive=io.BytesIO()
        with zipfile.ZipFile(archive,'w') as z:z.writestr('附件/图表.docx',original)
        _,source,assets=self.upload('资料包.zip',archive.getvalue())
        self.assertIn('附件/图表.docx',assets['items'][0]['source_path'])
        self.assertEqual(self.client.get(assets['items'][0]['source_url']).content,original)

    def test_failed_visual_extraction_retains_source_and_marks_failure(self):
        with patch('main.parse_source_file',return_value=main.ParseResult('正文仍可读取','pdf','parsed',None,'direct')):
            result,source,assets=self.upload('损坏图层.pdf',b'not pdf')
        self.assertEqual(source['parse_status'],'parsed')
        self.assertEqual(assets['status'],'failed')
        self.assertEqual(result['visual_assets']['incomplete_sources'],1)
        self.assertTrue(assets['error'])
        self.assertEqual(self.client.get(assets['source_url']).content,b'not pdf')
        self.assertEqual(self.client.get(assets['manifest_url']).json()['status'],'failed')


if __name__=='__main__':unittest.main()
