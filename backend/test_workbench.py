import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

import os
os.environ.setdefault('LLM_RETRIEVAL_ENABLED', 'false')
import main
from test_accounts_support import login, seed_demo


class WorkbenchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_patch = patch.object(main, "DB_PATH", Path(self.temp.name) / "test.db")
        self.db_patch.start()
        self.client = TestClient(main.app)
        self.client.__enter__()
        self.demo = seed_demo(self.client)
        self.run = self.demo["run"]

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.db_patch.stop()
        self.temp.cleanup()

    def test_demo_counts_categories_and_real_evidence(self):
        self.assertEqual(self.run["counts"], {"pass": 4, "fail": 0, "pending": 3, "not_applicable": 0, "total": 7})
        self.assertEqual(self.run["result"], "pending")
        self.assertEqual(self.run["topic_result"], "pending")
        self.assertEqual(sum(c["total"] for c in self.run["categories"]), 7)
        chunks = {c["id"]: c["text"] for c in self.run["chunks"]}
        for check in self.run["checks"]:
            for evidence in check["evidence"]:
                self.assertEqual(evidence["text"], chunks[evidence["id"]])

    def test_manual_qualification_retains_findings_and_persists_in_history(self):
        run_id = self.run['id']
        confirmed = self.client.post(f'/api/workbench/runs/{run_id}/qualify')
        self.assertEqual(confirmed.status_code, 200)
        saved = confirmed.json()
        self.assertEqual(saved['qualification']['result'], 'qualified')
        self.assertEqual(saved['qualification']['original_result'], 'pending')
        self.assertEqual(saved['status'], 'archived')
        for key in ('result', 'checks', 'counts', 'sources', 'chunks', 'quality_score'):
            self.assertEqual(saved[key], self.run[key])
        repeat = self.client.post(f'/api/workbench/runs/{run_id}/qualify').json()
        self.assertEqual(repeat['qualification'], saved['qualification'])
        history = self.client.get('/api/workbench/projects/DEMO-REVIEW/runs').json()['runs']
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]['qualification'], saved['qualification'])
        edit = self.client.patch(f'/api/workbench/runs/{run_id}/checks/DEMO-001',json={'result':'pass','reason':'复核','suggestion':'保留'})
        self.assertEqual(edit.status_code, 409)
        export = self.client.get(f'/api/workbench/runs/{run_id}/report.docx')
        text = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
        self.assertIn('人工确认结论：合格', text)
        self.assertIn('原始自动评审结论：待确认', text)
        fresh = self.client.post('/api/workbench/projects/DEMO-REVIEW/runs',json={}).json()
        self.assertNotIn('qualification', fresh)

    def test_rules_edit_and_rerun_keep_old_snapshot(self):
        rule = self.run["checks"][3]["rule"].copy()
        rule.update(enabled=False, standard="新的规则标准", content="新的规则内容")
        login(self.client, admin=True)
        saved = self.client.put(f"/api/workbench/rules/{rule['id']}", json=rule)
        login(self.client)
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json()["version"], 2)
        self.assertEqual(saved.json()["origin"], "custom")
        old = self.client.get(f"/api/workbench/runs/{self.run['id']}").json()
        self.assertEqual(old["checks"][3]["rule"]["version"], 1)
        fresh = self.client.post("/api/workbench/projects/DEMO-REVIEW/runs", json={}).json()
        self.assertNotEqual(fresh["id"], old["id"])
        self.assertEqual(fresh["counts"]["total"], 6)
        self.assertEqual(fresh["counts"]["fail"], 0)
        self.assertEqual(self.client.get("/api/workbench/projects/DEMO-REVIEW/runs/latest").json()["id"], fresh["id"])

    def test_no_topic_brief_is_pending_and_empty_category_is_pending(self):
        self.client.put("/api/workbench/projects/DEMO-REVIEW/brief", json={"text": ""})
        fresh = self.client.post("/api/workbench/projects/DEMO-REVIEW/runs", json={}).json()
        self.assertEqual(fresh["topic_result"], "pending")
        self.assertTrue(all(c["result"] == "pending" for c in fresh["checks"] if c["category"] == "topic"))
        login(self.client, admin=True)
        for check in fresh["checks"]:
            if check["category"] == "topic":
                rule = check["rule"].copy()
                rule["enabled"] = False
                self.client.put(f"/api/workbench/rules/{rule['id']}", json=rule)
        login(self.client)
        empty = self.client.post("/api/workbench/projects/DEMO-REVIEW/runs", json={}).json()
        category = next(c for c in empty["categories"] if c["id"] == "topic")
        self.assertEqual(category["total"], 0)
        self.assertEqual(category["result"], "pending")

    def test_review_updates_report_and_archive_preserves_verdict(self):
        for check in self.run["checks"]:
            response = self.client.patch(f"/api/workbench/runs/{self.run['id']}/checks/{check['id']}", json={"result": "pass", "reason": "专家逐项核验并确认材料符合本项标准。", "suggestion": "保留当前材料与核验记录。"})
            self.assertEqual(response.status_code, 200)
        run = response.json()
        self.assertEqual(run["result"], "pass")
        self.assertEqual(run["counts"]["pass"], 7)
        docx = self.client.get(f"/api/workbench/runs/{run['id']}/report.docx")
        self.assertEqual(docx.status_code, 200)
        doc = Document(io.BytesIO(docx.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        for heading in ["1. 评审基本信息", "2. 综合评审结论", "3. 评审摘要", "4. 三类评审结果", "5. 主要问题", "6. 评审依据", "7. 详细评审意见", "8. 最终建议"]:
            self.assertIn(heading, text)
        self.assertIn("综合结论：符合。课题要求符合性：符合。", text)
        self.assertIn(run["summary"], text)
        self.assertIn(run["checks"][0]["evidence"][0]["text"], text)
        self.assertEqual(len(doc.tables[0].rows), 4)
        self.client.post(f"/api/workbench/runs/{run['id']}/archive")
        rejected = self.client.patch(f"/api/workbench/runs/{run['id']}/checks/DEMO-001", json={"result": "fail", "reason": "再次修改", "suggestion": "重新提交"})
        self.assertEqual(rejected.status_code, 409)

    def test_professional_codes_and_document_rule_category(self):
        with main.db() as conn:
            conn.execute("INSERT INTO projects VALUES(?,?,?,?,?,?)", ("real", "测试汽机项目", "TURB", "技术改造", "uploaded", main.utc_now()))
            conn.execute('INSERT INTO project_owners VALUES(?,?)', ('real','user1'))
        rules = self.client.get("/api/workbench/rules?project_id=real").json()["rules"]
        self.assertTrue(any(r["id"].startswith("TUR-") for r in rules))
        docs = [r for r in rules if r["id"].startswith("TUR-DOC-")]
        self.assertTrue(docs)
        self.assertTrue(all(r["category"] == "common" for r in docs))

    def test_history_is_not_current_material_and_empty_package_not_pass(self):
        source = self.run["sources"][0]
        self.client.patch(f"/api/projects/DEMO-REVIEW/sources/{source['id']}/category", json={"category": "review_opinion"})
        material = self.client.get("/api/workbench/projects/DEMO-REVIEW/materials").json()
        self.assertEqual(material["chunks"], [])
        self.assertEqual(self.client.post("/api/workbench/projects/DEMO-REVIEW/runs", json={}).status_code, 409)
        package = self.client.post("/api/projects/DEMO-REVIEW/material-packages", json={"name": "空资料包"}).json()
        self.assertEqual(self.client.post("/api/workbench/projects/DEMO-REVIEW/runs", json={"material_package_id": package["id"]}).status_code, 409)

    def test_unparsed_sources_do_not_turn_missing_fields_into_failure(self):
        with main.db() as conn:
            conn.execute("UPDATE document_sources SET parse_status='ocr_required' WHERE project_id='DEMO-REVIEW'")
        fresh = self.client.post("/api/workbench/projects/DEMO-REVIEW/runs", json={}).json()
        self.assertEqual(fresh["counts"]["fail"], 0)
        self.assertGreater(fresh["counts"]["pending"], 0)

    def test_invalid_presence_rule_does_not_save(self):
        r = self.run["checks"][0]["rule"].copy()
        r["terms"] = [" "]
        login(self.client, admin=True)
        response = self.client.put(f"/api/workbench/rules/{r['id']}", json=r)
        self.assertEqual(response.status_code, 422)


if __name__ == "__main__":
    unittest.main()
